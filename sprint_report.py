#!/usr/bin/env python3
"""Create an executive sprint report and developer burndown slides from Excel."""

from __future__ import annotations

import argparse
import json
import logging
import math
import os
import re
import sys
import tempfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
import yaml

LOGGER = logging.getLogger("sprint_report")

REQUIRED_COLUMNS = (
    "Developer",
    "Task",
    "Sprint",
    "Status",
    "Assigned Date",
    "Start Date",
    "End Date",
    "Time Spent",
)
STATUS_ALIASES = {
    "todo": "To Do",
    "to do": "To Do",
    "in progress": "In-progress",
    "in-progress": "In-progress",
    "inprogress": "In-progress",
    "done": "Done",
    "blocked": "Blocked",
}
DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%m/%d/%Y %H:%M",
)
LOG_TOKEN = re.compile(
    r"(?P<date>\d{4}[-/]\d{1,2}[-/]\d{1,2})\s*"
    r"(?::|=|[-–—])\s*(?P<hours>\d+(?:\.\d+)?)\s*"
    r"(?P<unit>hours?|hrs?|h|minutes?|mins?|m)?",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class ReportConfig:
    input_file: Path
    template_file: Path
    output_file: Path
    sprint_start: date | None
    sprint_end: date | None
    sprint_name: str | None
    model: str
    temperature: float
    max_tokens: int
    timeout: int
    api_key_env: str
    api_base: str | None
    summary_layout_index: int
    developer_layout_index: int
    log_level: str


@dataclass
class SprintTask:
    developer: str
    task: str
    sprint: str
    status: str
    assigned_date: date | None
    start_date: date | None
    end_date: date | None
    effort_by_date: dict[date, float] = field(default_factory=dict)

    @property
    def logged_hours(self) -> float:
        return sum(self.effort_by_date.values())


@dataclass
class DeveloperBrief:
    """LLM (or deterministic) narrative for a single developer."""

    summary: str
    risk: str
    actions: list[str] = field(default_factory=list)


@dataclass
class TeamBrief:
    """Cross-developer executive narrative and risk flags for the summary slide."""

    narrative: str
    risks: list[str] = field(default_factory=list)


def _resolve_path(value: str, config_dir: Path) -> Path:
    path = Path(value).expanduser()
    return path if path.is_absolute() else (config_dir / path).resolve()


def _optional_date(value: Any, field_name: str) -> date | None:
    if value in (None, ""):
        return None
    parsed = parse_date(value)
    if parsed is None:
        raise ValueError(f"Invalid {field_name} date in configuration: {value!r}")
    return parsed


def load_config(config_path: Path) -> ReportConfig:
    """Load YAML settings and apply supported environment-variable overrides."""
    with config_path.open("r", encoding="utf-8") as stream:
        raw = yaml.safe_load(stream) or {}
    if not isinstance(raw, dict):
        raise ValueError("Configuration YAML must contain a mapping at its root.")

    files = raw.get("files", {})
    sprint = raw.get("sprint", {})
    llm = raw.get("llm", {})
    powerpoint = raw.get("powerpoint", {})
    config_dir = config_path.resolve().parent

    def setting(env_name: str, section: dict[str, Any], key: str, default: Any) -> Any:
        return os.environ.get(env_name, section.get(key, default))

    start = _optional_date(setting("SPRINT_START", sprint, "start_date", None), "sprint start")
    end = _optional_date(setting("SPRINT_END", sprint, "end_date", None), "sprint end")
    if (start is None) != (end is None):
        raise ValueError("Configure both sprint.start_date and sprint.end_date, or neither.")
    if start and end and start > end:
        raise ValueError("sprint.start_date must be on or before sprint.end_date.")

    config = ReportConfig(
        input_file=_resolve_path(
            str(setting("SPRINT_INPUT_FILE", files, "input_file", "templates/sprint_tasks.xlsx")),
            config_dir,
        ),
        template_file=_resolve_path(
            str(setting("SPRINT_TEMPLATE_FILE", files, "template_file", "templates/master_template.pptx")),
            config_dir,
        ),
        output_file=_resolve_path(
            str(setting("SPRINT_OUTPUT_FILE", files, "output_file", "sprint_report.pptx")),
            config_dir,
        ),
        sprint_start=start,
        sprint_end=end,
        sprint_name=setting("SPRINT_NAME", sprint, "name", None),
        model=str(setting("LLM_MODEL", llm, "model", "ollama/llama3.1")),
        temperature=float(setting("LLM_TEMPERATURE", llm, "temperature", 0.2)),
        max_tokens=int(setting("LLM_MAX_TOKENS", llm, "max_tokens", 300)),
        timeout=int(setting("LLM_TIMEOUT", llm, "timeout", 45)),
        api_key_env=str(setting("LLM_API_KEY_ENV", llm, "api_key_env", "OPENAI_API_KEY")),
        api_base=setting("LLM_API_BASE", llm, "api_base", None),
        summary_layout_index=int(
            setting("PPT_SUMMARY_LAYOUT_INDEX", powerpoint, "summary_layout_index", 6)
        ),
        developer_layout_index=int(
            setting("PPT_DEVELOPER_LAYOUT_INDEX", powerpoint, "developer_layout_index", 6)
        ),
        log_level=str(os.environ.get("LOG_LEVEL", raw.get("log_level", "INFO"))).upper(),
    )
    if not config.model.strip():
        raise ValueError("llm.model cannot be empty.")
    if config.max_tokens < 1 or config.timeout < 1:
        raise ValueError("LLM max_tokens and timeout must be positive integers.")
    return config


def parse_date(value: Any, excel_epoch: datetime | None = None) -> date | None:
    """Parse common Excel and text date values without raising on bad cells."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch=excel_epoch) if excel_epoch else from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (OverflowError, TypeError, ValueError):
            return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_hours(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return max(0.0, float(value))
    match = re.fullmatch(
        r"\s*(\d+(?:\.\d+)?)\s*(hours?|hrs?|h|minutes?|mins?|m)?\s*",
        str(value),
        re.IGNORECASE,
    )
    if not match:
        return None
    amount = float(match.group(1))
    unit = (match.group(2) or "h").lower()
    hours = amount / 60 if unit in {"m", "min", "mins", "minute", "minutes"} else amount
    return max(0.0, hours)


def parse_time_spent(value: Any) -> tuple[dict[date, float], list[str]]:
    """Parse JSON mappings or day-wise text into date-to-hours values."""
    if value is None or value == "":
        return {}, []
    errors: list[str] = []
    entries: list[tuple[Any, Any]] = []

    if isinstance(value, dict):
        entries = list(value.items())
    elif isinstance(value, (int, float)):
        return {}, ["numeric Time Spent has no associated date"]
    else:
        text = str(value).strip()
        try:
            decoded = json.loads(text)
        except json.JSONDecodeError:
            decoded = None
        if isinstance(decoded, dict):
            entries = list(decoded.items())
        elif isinstance(decoded, list):
            for item in decoded:
                if isinstance(item, dict):
                    entry_date = item.get("date") or item.get("day")
                    amount = item.get("hours", item.get("time"))
                    entries.append((entry_date, amount))
                else:
                    errors.append(f"unsupported JSON time-log entry: {item!r}")
        elif decoded is not None:
            errors.append("Time Spent JSON must be an object or list")
        else:
            matches = list(LOG_TOKEN.finditer(text))
            for match in matches:
                entries.append((match.group("date"), f"{match.group('hours')}{match.group('unit') or 'h'}"))
            remainder = LOG_TOKEN.sub("", text)
            remainder = re.sub(r"[\s,;|]+", "", remainder)
            if remainder:
                errors.append(f"unrecognized time-log text: {remainder[:60]!r}")

    totals: dict[date, float] = defaultdict(float)
    for raw_date, raw_hours in entries:
        parsed_day = parse_date(raw_date)
        hours = _parse_hours(raw_hours)
        if parsed_day is None:
            errors.append(f"invalid time-log date: {raw_date!r}")
        elif hours is None:
            errors.append(f"invalid time-log duration for {raw_date!r}: {raw_hours!r}")
        else:
            totals[parsed_day] += hours
    return dict(totals), errors


def _normalize_status(value: Any) -> str:
    status = re.sub(r"\s+", " ", str(value or "").strip()).lower()
    return STATUS_ALIASES.get(status, str(value or "Unknown").strip() or "Unknown")


def load_tasks(input_file: Path, sprint_name: str | None = None) -> list[SprintTask]:
    """Read the first worksheet, validate its headers, and clean task rows."""
    if not input_file.is_file():
        raise FileNotFoundError(f"Excel input file does not exist: {input_file}")
    workbook = load_workbook(input_file, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = {
            re.sub(r"\s+", " ", str(cell.value or "").strip()).casefold(): index
            for index, cell in enumerate(header_cells)
            if cell.value is not None
        }
        missing = [column for column in REQUIRED_COLUMNS if column.casefold() not in headers]
        if missing:
            raise ValueError(f"Excel file is missing required columns: {', '.join(missing)}")

        tasks: list[SprintTask] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = {
                name: row[headers[name.casefold()]] if headers[name.casefold()] < len(row) else None
                for name in REQUIRED_COLUMNS
            }
            developer = str(values["Developer"] or "").strip()
            task_title = str(values["Task"] or "").strip()
            sprint = str(values["Sprint"] or "").strip()
            if not developer or not task_title:
                LOGGER.warning("Skipping row %d: Developer or Task is blank.", row_number)
                continue
            if sprint_name and sprint.casefold() != sprint_name.casefold():
                continue
            status = _normalize_status(values["Status"])
            if status not in {"To Do", "In-progress", "Done", "Blocked"}:
                LOGGER.warning(
                    "Row %d (%s / %s) has an unrecognized status: %s",
                    row_number,
                    developer,
                    task_title,
                    status,
                )

            parsed_dates: dict[str, date | None] = {}
            for name in ("Assigned Date", "Start Date", "End Date"):
                parsed_dates[name] = parse_date(values[name], workbook.epoch)
                if values[name] not in (None, "") and parsed_dates[name] is None:
                    LOGGER.warning("Row %d has an invalid %s value: %r", row_number, name, values[name])

            effort, errors = parse_time_spent(values["Time Spent"])
            for error in errors:
                LOGGER.warning("Row %d (%s / %s): %s.", row_number, developer, task_title, error)
            tasks.append(
                SprintTask(
                    developer=developer,
                    task=task_title,
                    sprint=sprint or "Unspecified",
                    status=status,
                    assigned_date=parsed_dates["Assigned Date"],
                    start_date=parsed_dates["Start Date"],
                    end_date=parsed_dates["End Date"],
                    effort_by_date=effort,
                )
            )
        if not tasks:
            raise ValueError("No valid task rows were found in the workbook.")
        return tasks
    finally:
        workbook.close()


def resolve_sprint_range(
    tasks: Iterable[SprintTask], configured_start: date | None, configured_end: date | None
) -> tuple[date, date]:
    """Use configured dates or infer a useful inclusive range from task activity."""
    if configured_start and configured_end:
        return configured_start, configured_end
    observed = [
        day
        for task in tasks
        for day in (
            task.assigned_date,
            task.start_date,
            task.end_date,
            *task.effort_by_date.keys(),
        )
        if day is not None
    ]
    if not observed:
        raise ValueError("Cannot infer sprint dates: no valid task dates or time logs were found.")
    start = configured_start or min(observed)
    end = configured_end or max(observed)
    if start > end:
        raise ValueError("Resolved sprint start date is after sprint end date.")
    return start, end


def _task_in_range(task: SprintTask, start: date, end: date) -> bool:
    dates = [
        day
        for day in (task.assigned_date, task.start_date, task.end_date, *task.effort_by_date)
        if day is not None
    ]
    return not dates or min(dates) <= end and max(dates) >= start


def _completion_date(task: SprintTask, sprint_end: date) -> date:
    if task.end_date:
        return task.end_date
    logged_dates = [day for day in task.effort_by_date if day <= sprint_end]
    return max(logged_dates) if logged_dates else sprint_end


def burndown_series(
    tasks: list[SprintTask], start: date, end: date
) -> tuple[list[date], list[int], list[int], list[float]]:
    """Return daily ideal/actual remaining tasks and hours logged that day.

    Since the source workbook has no estimates, the burn-down unit is one task.
    """
    days = [start + timedelta(days=offset) for offset in range((end - start).days + 1)]
    total_tasks = len(tasks)
    completions: dict[date, int] = defaultdict(int)
    daily_hours: dict[date, float] = defaultdict(float)
    already_done = 0

    for task in tasks:
        if task.status == "Done":
            completed_on = _completion_date(task, end)
            if completed_on < start:
                already_done += 1
            elif completed_on <= end:
                completions[completed_on] += 1
        for logged_day, hours in task.effort_by_date.items():
            if start <= logged_day <= end:
                daily_hours[logged_day] += hours

    ideal = [
        max(0, math.ceil(total_tasks * (1 - (index + 1) / max(1, len(days)))))
        for index in range(len(days))
    ]
    completed_to_date = already_done
    actual: list[int] = []
    effort: list[float] = []
    for day in days:
        completed_to_date += completions[day]
        actual.append(max(0, total_tasks - completed_to_date))
        effort.append(daily_hours[day])
    return days, ideal, actual, effort


def _format_date(value: date | None) -> str:
    return value.strftime("%Y-%m-%d") if value else "—"


def _insight_line(tasks: list[SprintTask]) -> str:
    """Return a one-line factual highlight/blocker summary for a developer."""
    hours = sum(task.logged_hours for task in tasks)
    done = sum(task.status == "Done" for task in tasks)
    blocked = [task.task for task in tasks if task.status == "Blocked"]
    detail = f"Blockers: {', '.join(blocked[:4])}" if blocked else "No blocked tasks recorded"
    return f"Logged effort: {hours:.1f} h  |  Completed: {done}/{len(tasks)}  |  {detail}"


def _fallback_summary(tasks: list[SprintTask]) -> str:
    done = sum(task.status == "Done" for task in tasks)
    blocked = [task.task for task in tasks if task.status == "Blocked"]
    hours = sum(task.logged_hours for task in tasks)
    completion = done / len(tasks) * 100 if tasks else 0
    first = (
        f"{done} of {len(tasks)} tasks are complete ({completion:.0f}%), "
        f"with {hours:.1f} logged hours."
    )
    if blocked:
        second = f"Blocked work requires attention: {', '.join(blocked[:3])}."
    else:
        second = "No tasks are currently marked blocked; review in-progress work for delivery risks."
    return f"{first} {second}"


def _clean_sentences(text: str, max_sentences: int = 2) -> str:
    """Collapse whitespace and keep at most ``max_sentences`` sentences."""
    cleaned = re.sub(r"\s+", " ", text.strip())
    sentences = [sentence.strip() for sentence in re.split(r"(?<=[.!?])\s+", cleaned) if sentence.strip()]
    if not sentences:
        return ""
    result = " ".join(sentences[:max_sentences])
    if result[-1:] not in ".!?":
        result += "."
    return result


def _extract_json(text: str) -> dict[str, Any] | None:
    """Best-effort extraction of a JSON object from a model response."""
    if not text:
        return None
    cleaned = text.strip()
    if "```" in cleaned:
        # Drop fenced code markers and an optional ``json`` language hint.
        cleaned = re.sub(r"```(?:json)?", " ", cleaned, flags=re.IGNORECASE).strip()
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        parsed = json.loads(cleaned[start : end + 1])
    except json.JSONDecodeError:
        return None
    return parsed if isinstance(parsed, dict) else None


def _call_llm(system: str, prompt: str, config: ReportConfig) -> str | None:
    """Send one chat completion to the configured LiteLLM model.

    Returns the assistant text, or ``None`` when the provider is unreachable so
    callers can fall back to deterministic content.
    """
    try:
        import litellm
        from litellm import completion

        # Keep provider debug banners and retry chatter out of the application logs.
        litellm.suppress_debug_info = True
        kwargs: dict[str, Any] = {
            "model": config.model,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
            "temperature": config.temperature,
            "max_tokens": config.max_tokens,
            "timeout": config.timeout,
        }
        api_key = os.environ.get(config.api_key_env)
        if api_key:
            kwargs["api_key"] = api_key
        if config.api_base:
            kwargs["api_base"] = config.api_base
        response = completion(**kwargs)
        content = response.choices[0].message.content
        if isinstance(content, str) and content.strip():
            return content.strip()
        LOGGER.warning("LLM returned an empty response for model %s.", config.model)
    except Exception as exc:
        LOGGER.warning("LLM request failed (%s); using deterministic content.", exc)
    return None


def _fallback_brief(tasks: list[SprintTask]) -> DeveloperBrief:
    """Deterministic developer brief used when the LLM is unavailable."""
    summary = _fallback_summary(tasks)
    blocked = [task.task for task in tasks if task.status == "Blocked"]
    in_progress = [task.task for task in tasks if task.status == "In-progress"]
    todo = [task.task for task in tasks if task.status == "To Do"]
    risk = (
        f"{len(blocked)} blocked task(s) need attention: {', '.join(blocked[:3])}."
        if blocked
        else "No blocked tasks; monitor in-progress items for schedule risk."
    )
    actions: list[str] = []
    if blocked:
        actions.append(f"Unblock: {', '.join(blocked[:2])}")
    if in_progress:
        actions.append(f"Drive in-progress work to closure: {', '.join(in_progress[:2])}")
    if todo:
        actions.append(f"Start planned work: {', '.join(todo[:2])}")
    if not actions:
        actions.append("Confirm remaining scope and close out the sprint.")
    return DeveloperBrief(summary=summary, risk=risk, actions=actions[:3])


def generate_developer_brief(tasks: list[SprintTask], config: ReportConfig) -> DeveloperBrief:
    """Return a two-sentence summary, a risk sentence, and action items.

    The LLM is asked for strict JSON; any failure or malformed response falls
    back to a deterministic brief so the report is always produced.
    """
    if not tasks:
        return DeveloperBrief(
            summary=(
                "No task data is available for this developer. "
                "Please review the source workbook and sprint selection."
            ),
            risk="No data to assess.",
            actions=["Verify the workbook rows for this developer."],
        )
    fallback = _fallback_brief(tasks)
    records = [
        {
            "task": task.task,
            "status": task.status,
            "start": _format_date(task.start_date),
            "end": _format_date(task.end_date),
            "logged_hours": round(task.logged_hours, 2),
        }
        for task in tasks
    ]
    prompt = (
        "You are writing an executive sprint brief for one developer. Use only the "
        "provided facts; never invent causes, estimates, or results. Respond with "
        'STRICT JSON only, no prose, shaped as {"summary": string, "risk": string, '
        '"actions": [string, ...]}. "summary" must be exactly two concise sentences '
        "(a delivery highlight then an outcome). \"risk\" must be one sentence naming "
        "the most material blocker or risk, or state that none exist. \"actions\" must "
        "be 1-3 short imperative next steps. Developer task history: "
        + json.dumps(records, ensure_ascii=False)
    )
    raw = _call_llm("You write factual, concise sprint briefs as JSON.", prompt, config)
    data = _extract_json(raw) if raw else None
    if not data:
        return fallback
    summary = _clean_sentences(str(data.get("summary", "")), 2) or fallback.summary
    risk = str(data.get("risk", "")).strip() or fallback.risk
    raw_actions = data.get("actions")
    if isinstance(raw_actions, str):
        raw_actions = [raw_actions]
    actions = [str(item).strip() for item in (raw_actions or []) if str(item).strip()]
    return DeveloperBrief(summary=summary, risk=risk, actions=(actions or fallback.actions)[:3])


def generate_summary(tasks: list[SprintTask], config: ReportConfig) -> str:
    """Backwards-compatible helper returning only the two-sentence summary."""
    return generate_developer_brief(tasks, config).summary


def _fallback_team_brief(tasks: list[SprintTask]) -> TeamBrief:
    """Deterministic team narrative and cross-developer risk flags."""
    narrative = _fallback_summary(tasks)
    blocked_by_developer: dict[str, list[str]] = defaultdict(list)
    for task in tasks:
        if task.status == "Blocked":
            blocked_by_developer[task.developer].append(task.task)
    risks = [
        f"{developer}: {len(items)} blocked ({', '.join(items[:2])})"
        for developer, items in sorted(blocked_by_developer.items())
    ]
    if not risks:
        risks = ["No cross-team blockers; protect the sprint goal by clearing in-progress work."]
    return TeamBrief(narrative=narrative, risks=risks[:4])


def generate_team_brief(
    briefs: dict[str, DeveloperBrief], tasks: list[SprintTask], config: ReportConfig
) -> TeamBrief:
    """Return the team-level executive narrative and cross-developer risk flags."""
    if not tasks:
        return TeamBrief(
            narrative="No task data is available for the selected sprint window.",
            risks=["Verify the workbook and sprint selection."],
        )
    fallback = _fallback_team_brief(tasks)
    metrics = {
        "tasks": len(tasks),
        "done": sum(task.status == "Done" for task in tasks),
        "blocked": sum(task.status == "Blocked" for task in tasks),
        "developers": len({task.developer for task in tasks}),
        "logged_hours": round(sum(task.logged_hours for task in tasks), 1),
    }
    context = {
        "metrics": metrics,
        "developers": {
            name: {"summary": brief.summary, "risk": brief.risk, "actions": brief.actions}
            for name, brief in sorted(briefs.items())
        },
    }
    prompt = (
        "You are writing the executive summary for a team sprint report. Use only the "
        'provided facts. Respond with STRICT JSON only, shaped as {"narrative": string, '
        '"risks": [string, ...]}. "narrative" must be 2-3 concise sentences summarising '
        "overall delivery and outlook. \"risks\" must be 1-4 short cross-developer risk "
        "flags, naming the developer where relevant. Team data: "
        + json.dumps(context, ensure_ascii=False)
    )
    raw = _call_llm("You write factual, concise executive sprint summaries as JSON.", prompt, config)
    data = _extract_json(raw) if raw else None
    if not data:
        return fallback
    narrative = _clean_sentences(str(data.get("narrative", "")), 3) or fallback.narrative
    raw_risks = data.get("risks")
    if isinstance(raw_risks, str):
        raw_risks = [raw_risks]
    risks = [str(item).strip() for item in (raw_risks or []) if str(item).strip()]
    return TeamBrief(narrative=narrative, risks=(risks or fallback.risks)[:4])


def _chart_for_developer(
    developer: str, tasks: list[SprintTask], start: date, end: date, output: Path
) -> None:
    days, ideal, actual, effort = burndown_series(tasks, start, end)
    fig, axis = plt.subplots(figsize=(10, 4.4), dpi=220)
    axis.plot(days, ideal, color="#738096", linestyle="--", linewidth=2, label="Ideal remaining tasks")
    axis.step(days, actual, where="post", color="#176B87", linewidth=2.5, label="Actual remaining tasks")
    axis.fill_between(days, actual, step="post", color="#176B87", alpha=0.09)
    axis.set_ylabel("Remaining tasks", color="#243247")
    axis.set_ylim(bottom=0, top=max(1, len(tasks) * 1.12))
    axis.grid(axis="y", color="#DDE3EA", linewidth=0.8)
    axis.set_axisbelow(True)
    axis.xaxis.set_major_formatter(mdates.DateFormatter("%b %d"))
    axis.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=4, maxticks=8))
    axis.tick_params(axis="x", labelrotation=25, colors="#455468")
    axis.tick_params(axis="y", colors="#455468")

    effort_axis = axis.twinx()
    effort_axis.bar(days, effort, width=0.65, color="#E5A53B", alpha=0.62, label="Hours logged")
    effort_axis.set_ylabel("Hours logged per day", color="#966616")
    effort_axis.tick_params(axis="y", colors="#966616")
    handles_a, labels_a = axis.get_legend_handles_labels()
    handles_b, labels_b = effort_axis.get_legend_handles_labels()
    axis.legend(handles_a + handles_b, labels_a + labels_b, loc="upper right", frameon=False, ncols=2)
    axis.set_title(f"{developer} — Task Burndown and Daily Logged Effort", loc="left", fontweight="bold")
    fig.tight_layout()
    fig.savefig(output, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def _select_layout(presentation: Presentation, requested_index: int):
    layouts = presentation.slide_layouts
    if 0 <= requested_index < len(layouts):
        return layouts[requested_index]
    for layout in layouts:
        if "blank" in layout.name.casefold():
            LOGGER.warning(
                "Slide layout index %d is unavailable; using blank layout %r.",
                requested_index,
                layout.name,
            )
            return layout
    LOGGER.warning(
        "Slide layout index %d is unavailable and no blank layout exists; using first layout.",
        requested_index,
    )
    return layouts[0]


def _add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    font_size: float = 12,
    bold: bool = False,
    color: tuple[int, int, int] = (36, 50, 71),
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)
    return box


def _add_title(slide, title: str, width: float) -> None:
    title_shape = next(
        (
            shape
            for shape in slide.placeholders
            if getattr(shape.placeholder_format, "type", None) == 1
        ),
        None,
    )
    if title_shape is not None:
        title_shape.text = title
        for paragraph in title_shape.text_frame.paragraphs:
            paragraph.font.name = "Aptos Display"
            paragraph.font.size = Pt(25)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(23, 54, 76)
    else:
        _add_text(slide, title, 0.45, 0.22, width - 0.9, 0.48, font_size=25, bold=True, color=(23, 54, 76))


def _add_stat_card(slide, title: str, value: str, left: float, top: float, width: float) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(1.12),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(243, 247, 250)
    shape.line.color.rgb = RGBColor(220, 228, 235)
    shape.text_frame.clear()
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.text_frame.paragraphs[0].text = value
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].font.size = Pt(24)
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(23, 107, 135)
    paragraph = shape.text_frame.add_paragraph()
    paragraph.text = title
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(10)
    paragraph.font.color.rgb = RGBColor(69, 84, 104)


def _add_summary_slide(
    presentation: Presentation,
    config: ReportConfig,
    tasks: list[SprintTask],
    start: date,
    end: date,
    team_brief: TeamBrief,
):
    slide = presentation.slides.add_slide(
        _select_layout(presentation, config.summary_layout_index)
    )
    width = presentation.slide_width / Inches(1)
    _add_title(slide, "Sprint Executive Summary", width)
    total = len(tasks)
    done = sum(task.status == "Done" for task in tasks)
    blocked = sum(task.status == "Blocked" for task in tasks)
    developers = len({task.developer for task in tasks})
    total_hours = sum(task.logged_hours for task in tasks)
    card_width = (width - 1.2) / 4
    values = (
        ("Tasks completed", f"{done}/{total}"),
        ("Completion rate", f"{(done / total * 100 if total else 0):.0f}%"),
        ("Blocked tasks", str(blocked)),
        ("Logged effort", f"{total_hours:.1f} h"),
    )
    for index, (label, value) in enumerate(values):
        _add_stat_card(slide, label, value, 0.45 + index * (card_width + 0.1), 1.05, card_width)
    sprint_text = f"Sprint: {config.sprint_name or 'All selected tasks'}  |  {start:%b %d, %Y} – {end:%b %d, %Y}"
    _add_text(slide, sprint_text, 0.5, 2.32, width - 1.0, 0.35, font_size=15, bold=True)
    _add_text(slide, team_brief.narrative, 0.5, 2.72, width - 1.0, 1.0, font_size=14)
    _add_text(
        slide,
        "Risks & attention",
        0.5,
        3.80,
        width - 1.0,
        0.28,
        font_size=12,
        bold=True,
        color=(150, 80, 55),
    )
    _add_text(
        slide,
        "  •  ".join(team_brief.risks),
        0.5,
        4.08,
        width - 1.0,
        0.7,
        font_size=11,
        color=(69, 84, 104),
    )
    by_status: dict[str, int] = defaultdict(int)
    for task in tasks:
        by_status[task.status] += 1
    status_line = "  •  ".join(f"{status}: {count}" for status, count in sorted(by_status.items()))
    _add_text(slide, status_line, 0.5, 4.95, width - 1.0, 0.45, font_size=13, color=(69, 84, 104))
    _add_text(
        slide,
        "Developer slides include task-level status, dates, logged hours, an individual daily burndown, "
        "a risk note, and suggested next actions.",
        0.5,
        5.55,
        width - 1.0,
        0.55,
        font_size=12,
        color=(69, 84, 104),
    )
    return slide


def _add_task_table(slide, tasks: list[SprintTask], left: float, top: float, width: float, height: float) -> None:
    columns = ("Task", "Sprint", "Status", "Start", "End")
    row_count = len(tasks) + 1
    table_shape = slide.shapes.add_table(
        row_count, len(columns), Inches(left), Inches(top), Inches(width), Inches(height)
    )
    table = table_shape.table
    fractions = (0.39, 0.16, 0.17, 0.14, 0.14)
    for column, fraction in enumerate(fractions):
        table.columns[column].width = Inches(width * fraction)
    row_height = height / row_count
    font_size = max(5.5, min(10.5, 10.5 * min(1, 14 / max(1, row_count - 1))))
    for row_index, row in enumerate(table.rows):
        row.height = Inches(row_height)
        for column_index, cell in enumerate(row.cells):
            cell.margin_left = Inches(0.045)
            cell.margin_right = Inches(0.035)
            cell.margin_top = Inches(0.015)
            cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                RGBColor(23, 54, 76) if row_index == 0 else
                RGBColor(246, 248, 250) if row_index % 2 else RGBColor(255, 255, 255)
            )
            cell.text = columns[column_index] if row_index == 0 else ""
            if row_index:
                task = tasks[row_index - 1]
                values = (
                    task.task,
                    task.sprint,
                    task.status,
                    _format_date(task.start_date),
                    _format_date(task.end_date),
                )
                cell.text = values[column_index]
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(font_size)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(36, 50, 71)


def _add_developer_slide(
    presentation: Presentation,
    config: ReportConfig,
    developer: str,
    tasks: list[SprintTask],
    brief: DeveloperBrief,
    start: date,
    end: date,
    chart_path: Path,
):
    slide = presentation.slides.add_slide(
        _select_layout(presentation, config.developer_layout_index)
    )
    width = presentation.slide_width / Inches(1)
    height = presentation.slide_height / Inches(1)
    _add_title(slide, developer, width)
    _add_text(slide, brief.summary, 0.48, 0.80, width - 0.96, 0.60, font_size=13)
    _add_text(
        slide,
        _insight_line(tasks),
        0.5,
        1.42,
        width - 1.0,
        0.26,
        font_size=10,
        color=(69, 84, 104),
    )
    _add_text(slide, f"Risk: {brief.risk}", 0.5, 1.70, width - 1.0, 0.26, font_size=10, color=(150, 80, 55))
    _add_text(
        slide,
        "Next: " + "; ".join(brief.actions) if brief.actions else "Next: —",
        0.5,
        1.98,
        width - 1.0,
        0.34,
        font_size=10,
        color=(23, 107, 135),
    )

    content_top = 2.5
    content_height = max(2.0, height - content_top - 0.48)
    table_width = width * 0.51
    _add_task_table(slide, tasks, 0.45, content_top, table_width, content_height)
    _add_text(
        slide,
        f"Task history ({len(tasks)} tasks)",
        0.48,
        content_top - 0.28,
        table_width,
        0.25,
        font_size=10,
        bold=True,
    )
    chart_left = table_width + 0.72
    chart_width = width - chart_left - 0.4
    slide.shapes.add_picture(
        str(chart_path),
        Inches(chart_left),
        Inches(content_top + 0.12),
        width=Inches(chart_width),
        height=Inches(max(1.75, content_height - 0.2)),
    )
    return slide


def create_presentation(
    config: ReportConfig,
    tasks: list[SprintTask],
    start: date,
    end: date,
    briefs: dict[str, DeveloperBrief] | None = None,
    team_brief: TeamBrief | None = None,
) -> Path:
    """Build the presentation while retaining all existing template slides.

    ``briefs`` and ``team_brief`` may be supplied by a caller (for example the
    Streamlit UI) to avoid regenerating LLM content twice; otherwise they are
    produced here.
    """
    if not config.template_file.is_file():
        raise FileNotFoundError(f"PowerPoint template does not exist: {config.template_file}")
    protected_paths = {
        config.input_file.resolve(),
        config.template_file.resolve(),
    }
    if config.output_file.resolve() in protected_paths:
        raise ValueError("Output PowerPoint path must differ from the Excel input and template.")

    by_developer: dict[str, list[SprintTask]] = defaultdict(list)
    for task in tasks:
        by_developer[task.developer].append(task)
    ordered = sorted(by_developer.items())
    if briefs is None:
        briefs = {developer: generate_developer_brief(items, config) for developer, items in ordered}
    if team_brief is None:
        team_brief = generate_team_brief(briefs, tasks, config)

    presentation = Presentation(config.template_file)
    generated_slides = [_add_summary_slide(presentation, config, tasks, start, end, team_brief)]

    config.output_file.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="sprint-charts-") as temp_dir:
        for index, (developer, developer_tasks) in enumerate(ordered):
            chart = Path(temp_dir) / f"developer-{index}.png"
            _chart_for_developer(developer, developer_tasks, start, end, chart)
            generated_slides.append(
                _add_developer_slide(
                    presentation,
                    config,
                    developer,
                    developer_tasks,
                    briefs[developer],
                    start,
                    end,
                    chart,
                )
            )

        slide_id_list = presentation.slides._sldIdLst
        for slide in reversed(generated_slides):
            slide_id = next(
                item for item in slide_id_list if int(item.get("id")) == slide.slide_id
            )
            slide_id_list.remove(slide_id)
            slide_id_list.insert(0, slide_id)
        presentation.save(config.output_file)
    return config.output_file


def run(config_path: Path) -> Path:
    config = load_config(config_path)
    logging.basicConfig(
        level=getattr(logging, config.log_level, logging.INFO),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    # LiteLLM is chatty about provider selection; surface only warnings and above.
    logging.getLogger("LiteLLM").setLevel(logging.WARNING)
    tasks = load_tasks(config.input_file, config.sprint_name)
    start, end = resolve_sprint_range(tasks, config.sprint_start, config.sprint_end)
    tasks = [task for task in tasks if _task_in_range(task, start, end)]
    if not tasks:
        raise ValueError("No tasks overlap the selected sprint date range.")
    output = create_presentation(config, tasks, start, end)
    LOGGER.info("Created sprint report: %s", output)
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate sprint performance slides from an Excel workbook."
    )
    parser.add_argument(
        "--config",
        type=Path,
        default=Path("config.yaml"),
        help="YAML configuration path (default: ./config.yaml)",
    )
    args = parser.parse_args(argv)
    try:
        run(args.config)
        return 0
    except Exception as exc:
        logging.basicConfig(level=logging.ERROR, format="%(levelname)s: %(message)s")
        LOGGER.error("Report generation failed: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
